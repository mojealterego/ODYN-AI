from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


class _OdynPDF(FPDF):
    def __init__(self) -> None:
        super().__init__()
        self.set_compression(False)
        self.alias_nb_pages()
        self.set_margins(18, 20, 18)
        self.set_auto_page_break(auto=True, margin=20)

    def footer(self) -> None:
        self.set_y(-13)
        self.set_font("ODYN", "", 8)
        self.set_text_color(120, 130, 140)
        self.cell(0, 6, text=f"ODYN AI  |  {self.page_no()} / {{nb}}", align="C")


class OdynDocumentBuilder:
    """Native, professional export of ODYN results to PDF, DOCX and XLSX."""

    _HEADING_RE = re.compile(r"^(#{1,3})\s+(.+?)\s*$")

    def __init__(self, output_dir: str = "exports") -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _path(self, filename: str, extension: str) -> Path:
        safe_name = Path(filename).name.strip()
        if safe_name.lower().endswith(extension):
            safe_name = safe_name[: -len(extension)]
        if not safe_name:
            safe_name = "raport-odyn"
        return self.output_dir / f"{safe_name}{extension}"

    def _find_unicode_font(self) -> str | None:
        candidates = (
            os.getenv("ODYN_PDF_FONT"),
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        )
        for candidate in candidates:
            if candidate and Path(candidate).is_file():
                return candidate
        return None

    @staticmethod
    def _safe_xlsx_value(value: object) -> object:
        """Prevent spreadsheet formula injection while preserving real numbers."""
        if isinstance(value, str):
            stripped = value.lstrip()
            if stripped and stripped[0] in {"=", "+", "-", "@"}:
                return "'" + value
        return value

    def generate_pdf(self, title: str, content: str, filename: str) -> str:
        filepath = self._path(filename, ".pdf")
        pdf = _OdynPDF()
        font = self._find_unicode_font()
        if font:
            pdf.add_font("ODYN", "", font)
            pdf.add_font("ODYN", "B", font)
        else:
            pdf.add_font("ODYN", "", "Helvetica")
            pdf.add_font("ODYN", "B", "Helvetica")

        pdf.add_page()
        pdf.set_font("ODYN", "B", 17)
        pdf.set_text_color(230, 233, 236)
        pdf.cell(0, 10, text=title, new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.set_draw_color(184, 138, 50)
        pdf.line(18, pdf.get_y() + 2, 192, pdf.get_y() + 2)
        pdf.ln(8)
        pdf.set_font("ODYN", "", 10.5)
        pdf.set_text_color(35, 40, 45)

        for line in content.splitlines() or [""]:
            match = self._HEADING_RE.match(line)
            if match:
                level = len(match.group(1))
                pdf.ln(3)
                pdf.set_font("ODYN", "B", {1: 14, 2: 12, 3: 11}[level])
                pdf.set_text_color(70, 80, 90)
                pdf.multi_cell(0, 7, text=match.group(2))
                pdf.set_font("ODYN", "", 10.5)
                pdf.set_text_color(35, 40, 45)
            elif line.strip():
                pdf.multi_cell(0, 6.5, text=line)
            else:
                pdf.ln(3)

        pdf.output(str(filepath))
        return str(filepath)

    def generate_docx(self, title: str, content: str, filename: str) -> str:
        filepath = self._path(filename, ".docx")
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        styles = doc.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(10.5)
        for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
            styles[style_name].font.name = "Aptos Display"
            styles[style_name].font.size = Pt(size)
            styles[style_name].font.bold = True

        heading = doc.add_paragraph(style="Title")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run(title)

        header = section.header.paragraphs[0]
        header.text = "ODYN AI  ·  RAPORT"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(8)
        footer = section.footer.paragraphs[0]
        footer.text = "ODYN AI  ·  MojeAlterego"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.size = Pt(8)

        for line in content.splitlines():
            match = self._HEADING_RE.match(line)
            if match:
                doc.add_heading(match.group(2), level=len(match.group(1)))
            elif line.strip():
                doc.add_paragraph(line)
            else:
                doc.add_paragraph("")

        doc.save(filepath)
        return str(filepath)

    def generate_xlsx(
        self,
        data: Sequence[Sequence[object]],
        filename: str,
        *,
        sheet_name: str = "ODYN",
        header: bool = True,
    ) -> str:
        filepath = self._path(filename, ".xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31] or "ODYN"

        for row in data:
            ws.append([self._safe_xlsx_value(value) for value in row])

        if header and ws.max_row:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            if ws.max_row > 1:
                table_ref = f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}"
                table = Table(displayName="ODYNTable", ref=table_ref)
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                ws.add_table(table)
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = table_ref

        for column in ws.columns:
            values = [str(cell.value or "") for cell in column]
            width = min(max(len(value) for value in values) + 2, 60)
            ws.column_dimensions[column[0].column_letter].width = max(width, 10)

        ws.sheet_view.showGridLines = False
        wb.save(filepath)
        return str(filepath)

    def generate_report(
        self,
        format: str,
        title: str,
        content: str,
        filename: str,
        *,
        data: Sequence[Sequence[object]] | None = None,
        sheet_name: str = "ODYN",
        header: bool = True,
    ) -> str:
        """Route a generated agent report to the selected native format."""
        normalized = format.lower().lstrip(".")
        if normalized == "pdf":
            return self.generate_pdf(title, content, filename)
        if normalized == "docx":
            return self.generate_docx(title, content, filename)
        if normalized == "xlsx":
            if data is None:
                data = self._content_to_rows(content)
            return self.generate_xlsx(data, filename, sheet_name=sheet_name, header=header)
        raise ValueError("Nieobsługiwany format eksportu. Wybierz PDF, DOCX albo XLSX.")

    @staticmethod
    def _content_to_rows(content: str) -> list[list[str]]:
        rows: list[list[str]] = []
        for line in content.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if "|" in stripped:
                cells = [cell.strip() for cell in stripped.strip("|").split("|")]
                if any(cells):
                    rows.append(cells)
            else:
                rows.append([stripped])
        return rows or [["Brak danych"]]
