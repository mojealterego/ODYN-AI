import tempfile
import unittest
from pathlib import Path

from odyn_ai.core.document_generator import OdynDocumentBuilder


class DocumentGeneratorTests(unittest.TestCase):
    def test_generates_pdf_with_polish_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = OdynDocumentBuilder(tmp).generate_pdf(
                "Raport ODYN",
                "Zażółć gęślą jaźń. Pamięć epizodyczna.",
                "raport",
            )
            output = Path(path)
            self.assertTrue(output.is_file())
            self.assertGreater(output.stat().st_size, 100)

    def test_generates_docx_with_structure(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = OdynDocumentBuilder(tmp).generate_docx(
                "Raport ODYN",
                "Pierwszy akapit.\n\nDrugi akapit.",
                "raport",
            )
            from docx import Document
            doc = Document(path)
            self.assertEqual(doc.paragraphs[0].text, "Raport ODYN")
            self.assertIn("Pierwszy akapit.", [p.text for p in doc.paragraphs])
            self.assertIn("Drugi akapit.", [p.text for p in doc.paragraphs])

    def test_generates_formatted_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = OdynDocumentBuilder(tmp).generate_xlsx(
                [["Nazwa", "Wartość"], ["ODYN", 42]],
                "dane",
            )
            from openpyxl import load_workbook
            wb = load_workbook(path)
            ws = wb["ODYN"]
            self.assertEqual(ws["A1"].value, "Nazwa")
            self.assertTrue(ws["A1"].font.bold)
            self.assertEqual(ws["B2"].value, 42)
            self.assertEqual(ws.freeze_panes, "A2")


if __name__ == "__main__":
    unittest.main()
